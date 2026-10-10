import sys
import os
import time
import subprocess
import openpyxl
from PyQt6.QtWidgets import (QApplication, QMainWindow, QLabel, QVBoxLayout, 
                             QHBoxLayout, QWidget, QPushButton, QTextBrowser, QGraphicsDropShadowEffect)
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QPixmap, QColor

class AnalysisPage5(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("OncoVue - Patient Report Summary")
        self.setFixedSize(1280, 720)
        
        self.data_file = r"data/data.xlsx"
        self.user_file = r"data/user.xlsx"
        self.report_dir = r"data/report"
        self.phone_arg = sys.argv[1] if len(sys.argv) > 1 else ""
        
        self.patient_data = None
        self.doctor_address = "Consultant Clinical Diagnostic Department"
        
        os.makedirs(self.report_dir, exist_ok=True)
        self.load_patient_record()
        self.get_doctor_meta()
        
        # Central Widget & Background Setup (Soft Pastel Gradient Base)
        self.central_widget = QWidget()
        self.central_widget.setObjectName("CentralWidget")
        self.central_widget.setStyleSheet("""
            #CentralWidget {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:1, 
                    stop:0 #FAF7FF, 
                    stop:0.6 #F3E8FF, 
                    stop:1 #E9D5FF);
            }
        """)
        self.setCentralWidget(self.central_widget)
        
        # Master Vertical Layout
        self.main_layout = QVBoxLayout(self.central_widget)
        self.main_layout.setContentsMargins(25, 20, 25, 20)
        self.main_layout.setSpacing(15)
        
        # Top Navigation Bar Layout (Glass Back Button)
        self.top_bar = QHBoxLayout()
        self.btn_back = QPushButton("← Back")
        self.btn_back.setFixedSize(110, 38)
        self.btn_back.setStyleSheet("""
            QPushButton {
                background-color: rgba(109, 40, 217, 0.08);
                border: 1px solid rgba(109, 40, 217, 0.18);
                color: #4C1D95; /* Velvet Slate */
                font-family: 'Poppins', 'Inter', sans-serif;
                font-weight: bold;
                border-radius: 8px;
                font-size: 13px;
            }
            QPushButton:hover {
                background-color: rgba(109, 40, 217, 0.15);
                border-color: rgba(109, 40, 217, 0.3);
            }
            QPushButton:pressed {
                background-color: rgba(109, 40, 217, 0.25);
            }
        """)
        self.btn_back.clicked.connect(self.back_to_page4)
        self.top_bar.addWidget(self.btn_back)
        self.top_bar.addStretch()
        self.main_layout.addLayout(self.top_bar)
        
        # Frosted glass report preview browser
        self.report_view = QTextBrowser()
        self.report_view.setStyleSheet("""
            QTextBrowser {
                background-color: rgba(255, 255, 255, 0.85);
                border: 1px solid rgba(109, 40, 217, 0.15);
                border-radius: 12px;
                padding: 20px;
                font-family: 'Poppins', 'Inter', sans-serif;
                font-size: 13px;
                color: #2E1065; /* Deep Amethyst */
            }
        """)
        self.main_layout.addWidget(self.report_view)
        
        # Bottom Navigation Controls layout
        self.bottom_bar = QHBoxLayout()
        self.bottom_bar.setSpacing(15)
        self.main_layout.addLayout(self.bottom_bar)
        
        btn_style = """
            QPushButton {
                font-family: 'Poppins', 'Inter', sans-serif;
                font-weight: bold;
                font-size: 14px;
                border: none;
                border-radius: 8px;
                padding: 11px 24px;
                color: white;
            }
            QPushButton:disabled {
                background-color: rgba(109, 40, 217, 0.08);
                color: rgba(109, 40, 217, 0.35);
                border: 1px solid rgba(109, 40, 217, 0.05);
            }
        """
        
        # Primary Action Trigger (Rose-to-Orchid Gradient)
        self.btn_pdf = QPushButton("Generate PDF Report")
        self.btn_pdf.setStyleSheet(btn_style + """
            QPushButton {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0, 
                    stop:0 #EC4899, 
                    stop:1 #6D28D9);
            }
            QPushButton:hover {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0, 
                    stop:0 #F472B6, 
                    stop:1 #7C3AED);
            }
        """)
        self.btn_pdf.clicked.connect(self.generate_pdf_report)
        self.bottom_bar.addWidget(self.btn_pdf)
        
        # Secondary Action Trigger (Solid Orchid Purple Action)
        self.btn_see_report = QPushButton("See Report Directory")
        self.btn_see_report.setEnabled(False)
        self.btn_see_report.setStyleSheet(btn_style + """
            QPushButton {
                background-color: #6D28D9;
            }
            QPushButton:hover {
                background-color: #7C3AED;
            }
        """)
        self.btn_see_report.clicked.connect(self.see_report_page2)
        self.bottom_bar.addWidget(self.btn_see_report)
        
        # Operation Status Label
        self.status_lbl = QLabel("")
        self.status_lbl.setStyleSheet("""
            font-family: 'Poppins', sans-serif;
            color: #EF4444; /* Soft Crimson error warning */
            font-weight: bold; 
            font-size: 12px;
            background: transparent;
        """)
        self.main_layout.addWidget(self.status_lbl, alignment=Qt.AlignmentFlag.AlignCenter)
        
        # Apply shadow to the main preview container
        shadow_view = QGraphicsDropShadowEffect(self)
        shadow_view.setBlurRadius(25)
        shadow_view.setColor(QColor(109, 40, 217, 15))
        shadow_view.setOffset(0, 8)
        self.report_view.setGraphicsEffect(shadow_view)
        
        self.display_report_preview()

    def back_to_page4(self):
                # self.timer.stop()
        subprocess.Popen([sys.executable, "loader.py"])
        subprocess.Popen([sys.executable, "analysis_page4.py"])
        self.close()

    def load_patient_record(self):
        if not os.path.exists(self.data_file) or not self.phone_arg:
            return
        try:
            wb = openpyxl.load_workbook(self.data_file, data_only=True)
            sheet = wb.active
            headers = [str(cell.value).strip().lower() for cell in sheet[1]]
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if any(x is not None for x in row):
                    row_dict = {headers[i]: row[i] for i in range(len(headers)) if i < len(row)}
                    if str(row_dict.get("phone_no", "")).strip() == self.phone_arg:
                        self.patient_data = row_dict
                        break
            wb.close()
        except Exception:
            pass

    def get_doctor_meta(self):
        if not self.patient_data:
            return
        doc_name = self.patient_data.get("doctor", "")
        if os.path.exists(self.user_file):
            try:
                wb = openpyxl.load_workbook(self.user_file, data_only=True)
                sheet = wb.active
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if len(row) >= 6 and str(row[1]).strip() == str(doc_name).strip():
                        self.doctor_address = str(row[5]).strip() if row[5] else "Consultant Clinical Diagnostic Department"
                        break
                wb.close()
            except Exception:
                pass

    def display_report_preview(self):
        if not self.patient_data:
            self.report_view.setText("Error: No patient record found matching the current argument.")
            return
            
        data = self.patient_data
        
        # Reformatted inline preview layout designed with brand visual guidelines
        preview = f"""
        <div style="font-family: 'Poppins', 'Inter', sans-serif; line-height: 1.6; color: #2e1065; padding: 10px;">
            <h2 style="text-align: center; color: #6d28d9; margin-bottom: 20px; font-weight: bold; letter-spacing: 1px;">Patient Clinical Report Preview</h2>
            <hr style="border: none; border-bottom: 1px solid rgba(109, 40, 217, 0.15); margin-bottom: 20px;">
            
            <h3 style="color: #6d28d9; font-size: 14px; text-transform: uppercase; letter-spacing: 1.5px; margin-bottom: 10px;">Header Information</h3>
            <p style="margin: 4px 0;"><b style="color: #4c1d95;">Date:</b> {data.get('date')} &nbsp;|&nbsp; <b style="color: #4c1d95;">Time:</b> {data.get('time')}</p>
            <p style="margin: 4px 0;"><b style="color: #4c1d95;">Name:</b> {data.get('patien_name')} &nbsp;|&nbsp; <b style="color: #4c1d95;">Age:</b> {data.get('age')} &nbsp;|&nbsp; <b style="color: #4c1d95;">Gender:</b> {data.get('gender')}</p>
            <p style="margin: 4px 0;"><b style="color: #4c1d95;">Phone Number:</b> {data.get('phone_no')}</p>
            
            <hr style="border: none; border-bottom: 1px solid rgba(109, 40, 217, 0.15); margin-top: 20px; margin-bottom: 20px;">
            
            <h3 style="color: #6d28d9; font-size: 14px; text-transform: uppercase; letter-spacing: 1.5px; margin-bottom: 10px;">Main Analysis Findings</h3>
            <p style="margin: 4px 0;"><b style="color: #4c1d95;">Image Scan Path:</b> {data.get('image_path')}</p>
            <p style="margin: 4px 0; font-weight: bold; color: #4c1d95;">Classification Scores:</p>
            <ul style="margin: 5px 0 10px 20px; padding: 0;">
                <li style="margin: 2px 0;">Her0 Score: {data.get('her0')}</li>
                <li style="margin: 2px 0;">Her1 Score: {data.get('her1')}</li>
                <li style="margin: 2px 0;">Her2 Score: {data.get('her2')}</li>
                <li style="margin: 2px 0;">Her3 Score: {data.get('her3')}</li>
            </ul>
            <p style="margin: 4px 0;"><b style="color: #4c1d95;">HER Final Score Class:</b> {str(data.get('her_final')).upper()}</p>
            <p style="margin: 10px 0;"><b style="color: #4c1d95;">Final Prediction Verdict:</b> <span style="color: #10b981; font-weight: bold; font-size: 14px;">{data.get('final_prediction')}</span></p>
            
            <hr style="border: none; border-bottom: 1px solid rgba(109, 40, 217, 0.15); margin-top: 20px; margin-bottom: 20px;">
            
            <h3 style="color: #6d28d9; font-size: 14px; text-transform: uppercase; letter-spacing: 1.5px; margin-bottom: 10px;">Footer & Metadata</h3>
            <p style="margin: 4px 0;"><b style="color: #4c1d95;">Assigned Doctor:</b> {data.get('doctor')}</p>
            <p style="margin: 4px 0;"><b style="color: #4c1d95;">Hospital Name:</b> {data.get('hospital')}</p>
            <p style="margin: 4px 0;"><b style="color: #4c1d95;">Hospital Location / Specialization Address:</b> {self.doctor_address}</p>
        </div>
        """
        self.report_view.setHtml(preview)

    def generate_pdf_report(self):
        self.status_lbl.clear()
        if not self.patient_data:
            self.status_lbl.setText("Error: Patient data invalid.")
            return
            
        data = self.patient_data
        pdf_name = f"Report_{data.get('patien_name')}_{data.get('phone_no')}.pdf".replace(" ", "_")
        pdf_path = os.path.join(self.report_dir, pdf_name)
        
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            
            doc = SimpleDocTemplate(pdf_path, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
            story = []
            
            styles = getSampleStyleSheet()
            
            title_style = ParagraphStyle(
                'TitleStyle',
                parent=styles['Heading1'],
                fontSize=22,
                spaceAfter=15,
                textColor='#34495e',
                alignment=1
            )
            
            section_style = ParagraphStyle(
                'SectionStyle',
                parent=styles['Heading2'],
                fontSize=16,
                spaceBefore=10,
                spaceAfter=5,
                textColor='#2c3e50'
            )
            
            body_style = ParagraphStyle(
                'BodyStyle',
                parent=styles['Normal'],
                fontSize=13,
                spaceAfter=6,
                textColor='#34495e'
            )

            body_style_1 = ParagraphStyle(
                'BodyStyle',
                parent=styles['Normal'],
                
                fontSize=14,
                spaceAfter=6,
                textColor='#34495e'
            )
            
            story.append(Paragraph("OncoVue Clinical Report", title_style))
            story.append(Spacer(1, 10))
            
            story.append(Paragraph("Patient Demographics and Informstion", section_style))
            story.append(Paragraph(f"<b>Date:</b> {data.get('date')} | <b>Time:</b> {data.get('time')}", body_style))
            story.append(Paragraph(f"<b>Patient Name:</b> {data.get('patien_name')}", body_style))
            story.append(Paragraph(f"<b>Age:</b> {data.get('age')} | <b>Gender:</b> {data.get('gender')}", body_style))
            story.append(Paragraph(f"<b>Phone Number:</b> {data.get('phone_no')}", body_style))
            story.append(Spacer(1, 10))
            
            story.append(Paragraph("Diagnostic IHC HER2Neu Image ", section_style))
            img_path = str(data.get("image_path", "")).strip()
            if img_path and os.path.exists(img_path):
                story.append(Image(img_path, width=240, height=180))
            else:
                story.append(Paragraph("Image attachment not found.", body_style))
            story.append(Spacer(1, 10))
            
            story.append(Paragraph("Classification & Score Breakdown", section_style))
            story.append(Paragraph(f"<b>Her0 Score:</b> {data.get('her0')}", body_style))
            story.append(Paragraph(f"<b>Her1 Score:</b> {data.get('her1')}", body_style))
            story.append(Paragraph(f"<b>Her2 Score:</b> {data.get('her2')}", body_style))
            story.append(Paragraph(f"<b>Her3 Score:</b> {data.get('her3')}", body_style))
            story.append(Paragraph(f"<b>HER Final Score Class:</b> {str(data.get('her_final')).upper()}", body_style_1))
            story.append(Paragraph(f"<b>Final Clinical Verdict:</b> {data.get('final_prediction')}", body_style_1))
            story.append(Spacer(1, 10))
            
            story.append(Paragraph("Authorized Institution Information", section_style))
            story.append(Paragraph(f"<b>Attending Doctor:</b> {data.get('doctor')}", body_style))
            story.append(Paragraph(f"<b>Hospital Institution:</b> {data.get('hospital')}", body_style))
            story.append(Paragraph(f"<b>Address:</b> {self.doctor_address}", body_style))
            
            doc.build(story)
            
        except Exception as e:
            self.status_lbl.setText(f"Error compiling PDF Document: {str(e)}")
            return
            
        try:
            wb = openpyxl.load_workbook(self.data_file)
            sheet = wb.active
            headers = [str(cell.value).strip().lower() for cell in sheet[1]]
            
            row_idx = 2
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if any(x is not None for x in row):
                    if str(row[5]).strip() == self.phone_arg:
                        sheet.cell(row=row_idx, column=17).value = pdf_path
                        break
                row_idx += 1
            wb.save(self.data_file)
            wb.close()
        except Exception as e:
            self.status_lbl.setText(f"Error appending PDF path to directory: {str(e)}")
            return
            
        self.status_lbl.setStyleSheet("color: #10B981; font-family: 'Poppins', sans-serif; font-weight: bold;")
        self.status_lbl.setText(f"Success: Report generated at {pdf_path}")
        self.btn_see_report.setEnabled(True)
        
        self.load_pdf_report(pdf_path)

    def load_pdf_report(self, filepath):
        try:
            import pypdf
            reader = pypdf.PdfReader(filepath)
            
            # Incorporate premium design tokens inline with the extracted HTML structure
            html_content = "<div style='font-family: Poppins, Inter, sans-serif; line-height: 1.6; color: #2e1065; padding: 5px;'>"
            for idx, page in enumerate(reader.pages):
                extracted = page.extract_text()
                if extracted:
                    formatted_text = extracted.replace("\n", "<br>")
                    html_content += f"<div style='background-color: rgba(255, 255, 255, 0.95); padding: 25px; margin-bottom: 20px; border: 1px solid rgba(109, 40, 217, 0.12); border-radius: 8px; box-shadow: 0 4px 12px rgba(109, 40, 217, 0.03);'>"
                    html_content += f"<p style='text-align: right; font-size: 11px; color: #c084fc; border-bottom: 1px solid rgba(109, 40, 217, 0.1); padding-bottom: 5px; margin-bottom: 15px;'>Page {idx+1} of {len(reader.pages)}</p>"
                    html_content += f"<div style='font-size: 13px; font-weight: normal; color: #2e1065;'>{formatted_text}</div>"
                    html_content += "</div>"
            html_content += "</div>"
            if len(reader.pages) > 0:
                self.report_view.setHtml(html_content)
        except Exception:
            pass

    def see_report_page2(self):
        if self.patient_data:
            phone_no = str(self.patient_data.get("phone_no", ""))
                    # self.timer.stop()
            subprocess.Popen([sys.executable, "loader.py"])
            subprocess.Popen([sys.executable, "report_page2.py", phone_no])
            self.close()

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = AnalysisPage5()
    window.showFullScreen()
    window.show()
    sys.exit(app.exec())