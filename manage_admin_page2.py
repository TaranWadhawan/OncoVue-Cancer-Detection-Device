import sys
# import os
import os
os.environ["QT_QPA_PLATFORM"] = "xcb"
import re
import subprocess
import openpyxl
from PyQt6.QtWidgets import (QApplication, QMainWindow, QLabel, QVBoxLayout, 
                             QHBoxLayout, QWidget, QLineEdit, QPushButton, 
                             QComboBox, QTableWidget, QTableWidgetItem, QTabWidget, QHeaderView, QGraphicsDropShadowEffect)
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QColor
from keyboard import VirtualKeyboard

class AdminDashboard(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("OncoVue - Database Manager")
        self.setFixedSize(1280, 720)
        
        self.user_file = r"data/user.xlsx"
        self.manager_file = r"data/manager.xlsx"
        
        self.user_states = {}
        self.manager_states = {}
        
        self.init_excel_files()
        
        # Central Widget & Background Setup (Soft Pastel Gradient Base)
        self.central_widget = QWidget()
        self.central_widget.setObjectName("CentralWidget")
        self.central_widget.setStyleSheet("""
            #CentralWidget {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:1, 
                    stop:0 #FAF7FF, 
                    stop:0.6 #F3E8FF, 
                    stop:1 #E9D5FF);
            }
        """)
        self.setCentralWidget(self.central_widget)
        
        # Master Vertical Layout
        self.main_layout = QVBoxLayout(self.central_widget)
        self.main_layout.setContentsMargins(25, 20, 25, 20)
        self.main_layout.setSpacing(15)
        
        # Top Navigation Bar Layout (Glass Back Button)
        self.top_bar = QHBoxLayout()
        self.btn_back = QPushButton("← Back")
        self.btn_back.setObjectName("btn_back")
        self.btn_back.setFixedSize(110, 38)
        self.btn_back.setStyleSheet("""
            QPushButton#btn_back {
                background-color: rgba(109, 40, 217, 0.08);
                border: 1px solid rgba(109, 40, 217, 0.18);
                color: #4C1D95; /* Velvet Slate */
                font-family: 'Poppins', 'Inter', sans-serif;
                font-weight: bold;
                border-radius: 8px;
                font-size: 13px;
            }
            QPushButton#btn_back:hover {
                background-color: rgba(109, 40, 217, 0.15);
                border-color: rgba(109, 40, 217, 0.3);
            }
            QPushButton#btn_back:pressed {
                background-color: rgba(109, 40, 217, 0.25);
            }
        """)
        self.btn_back.clicked.connect(self.back_to_page3)
        self.top_bar.addWidget(self.btn_back)
        self.top_bar.addStretch()
        self.main_layout.addLayout(self.top_bar)
        
        # Frosted glass structured Tab Widget
        self.tab_widget = QTabWidget()
        self.tab_widget.setStyleSheet("""
            QTabWidget::pane {
                border: 1px solid rgba(109, 40, 217, 0.15);
                background-color: rgba(255, 255, 255, 0.65);
                border-radius: 12px;
            }
            QTabBar::tab {
                background-color: rgba(109, 40, 217, 0.05);
                color: #4C1D95;
                padding: 10px 25px;
                font-family: 'Poppins', 'Inter', sans-serif;
                font-weight: bold;
                font-size: 13px;
                border-top-left-radius: 8px;
                border-top-right-radius: 8px;
                margin-right: 2px;
            }
            QTabBar::tab:selected {
                background-color: rgba(255, 255, 255, 0.85);
                border: 1px solid rgba(109, 40, 217, 0.15);
                border-bottom-color: transparent;
                color: #2E1065; /* Deep Amethyst */
            }
        """)
        
        self.init_user_tab()
        self.init_manager_tab()
        
        self.main_layout.addWidget(self.tab_widget)
        self.refresh_user_table()
        self.refresh_manager_table()
        self.keyboard = VirtualKeyboard(self)
        self.keyboard.show()
        self.keyboard.move(100, 480)

    def init_excel_files(self):
        if not os.path.exists(self.user_file):
            wb = openpyxl.Workbook()
            sheet = wb.active
            sheet.append(["username", "doctor_name", "password", "reg_no", "hospital", "specialization", "phone", "email"])
            wb.save(self.user_file)
            wb.close()
            
        if not os.path.exists(self.manager_file):
            wb = openpyxl.Workbook()
            sheet = wb.active
            sheet.append(["Username", "Password"])
            wb.save(self.manager_file)
            wb.close()

    def back_to_page3(self):
                    # self.timer.stop()
        subprocess.Popen([sys.executable, "loader.py"])
        subprocess.Popen([sys.executable, "page3.py"])
        self.close()

    def init_user_tab(self):
        self.user_tab = QWidget()
        layout = QHBoxLayout(self.user_tab)
        layout.setContentsMargins(15, 15, 15, 15)
        
        # Frosted Glass form card
        form_panel = QWidget()
        form_panel.setObjectName("UserForm")
        form_panel.setFixedWidth(340)
        form_panel.setStyleSheet("""
            #UserForm {
                background-color: rgba(255, 255, 255, 0.72);
                border: 1px solid rgba(109, 40, 217, 0.15);
                border-radius: 18px;
            }
        """)
        
        shadow_user_form = QGraphicsDropShadowEffect(self)
        shadow_user_form.setBlurRadius(25)
        shadow_user_form.setColor(QColor(109, 40, 217, 15))
        form_panel.setGraphicsEffect(shadow_user_form)
        
        form_layout = QVBoxLayout(form_panel)
        form_layout.setSpacing(6)
        form_layout.setContentsMargins(20, 20, 20, 20)
        
        self.u_username = QLineEdit()
        self.u_username.setPlaceholderText("Username (Unique)")
        self.u_doctor_name = QLineEdit()
        self.u_doctor_name.setPlaceholderText("Doctor Name")
        self.u_password = QLineEdit()
        self.u_password.setPlaceholderText("Password")
        self.u_reg_no = QLineEdit()
        self.u_reg_no.setPlaceholderText("Registration No. (Unique)")
        self.u_hospital = QLineEdit()
        self.u_hospital.setPlaceholderText("Hospital")
        
        self.u_specialization = QComboBox()
        self.u_specialization.addItems(["Gen.Physician", "Phathology", "MD", "MS", "MBBS", "Oncology", "Gynacology", "obsetraian"])
        
        self.u_phone = QLineEdit()
        self.u_phone.setPlaceholderText("Phone (10 Digits)")
        self.u_email = QLineEdit()
        self.u_email.setPlaceholderText("Email")
        
        input_style = """
            QLineEdit, QComboBox {
                border: 1px solid rgba(109, 40, 217, 0.18);
                border-radius: 6px;
                padding: 6px 6px;
                height : 14px;
                font-family: 'Poppins', 'Inter', sans-serif;
                font-size: 12px;
                background-color: rgba(255, 255, 255, 0.9);
                color: #2E1065;
            }
            QLineEdit:focus, QComboBox:focus {
                border: 1.5px solid #6D28D9;
                background-color: #ffffff;
            }
            QComboBox QAbstractItemView {
                background-color: #ffffff;
                color: #2E1065;
                selection-background-color: #C084FC;
            }
        """
        self.u_username.setStyleSheet(input_style)
        self.u_doctor_name.setStyleSheet(input_style)
        self.u_password.setStyleSheet(input_style)
        self.u_reg_no.setStyleSheet(input_style)
        self.u_hospital.setStyleSheet(input_style)
        self.u_specialization.setStyleSheet(input_style)
        self.u_phone.setStyleSheet(input_style)
        self.u_email.setStyleSheet(input_style)
        
        lbl_style = "font-family: 'Poppins', sans-serif; font-size: 11px; color: #4C1D95; font-weight: bold;"
        
        form_layout.addWidget(QLabel("Username:", styleSheet=lbl_style))
        form_layout.addWidget(self.u_username)
        form_layout.addWidget(QLabel("Doctor Name:", styleSheet=lbl_style))
        form_layout.addWidget(self.u_doctor_name)
        form_layout.addWidget(QLabel("Password:", styleSheet=lbl_style))
        form_layout.addWidget(self.u_password)
        form_layout.addWidget(QLabel("Reg No:", styleSheet=lbl_style))
        form_layout.addWidget(self.u_reg_no)
        form_layout.addWidget(QLabel("Hospital:", styleSheet=lbl_style))
        form_layout.addWidget(self.u_hospital)
        form_layout.addWidget(QLabel("Specialization:", styleSheet=lbl_style))
        form_layout.addWidget(self.u_specialization)
        form_layout.addWidget(QLabel("Phone:", styleSheet=lbl_style))
        form_layout.addWidget(self.u_phone)
        form_layout.addWidget(QLabel("Email:", styleSheet=lbl_style))
        form_layout.addWidget(self.u_email)
        
        btn_style = """
            QPushButton {
                font-family: 'Poppins', 'Inter', sans-serif;
                font-weight: bold;
                font-size: 13px;
                border: none;
                border-radius: 6px;
                padding: 8px;
                color: white;
            }
        """
        
        self.btn_add_user = QPushButton("Add User")
        self.btn_add_user.setStyleSheet(btn_style + "QPushButton{background-color: #10B981;} QPushButton:hover{background-color: #059669;}")
        self.btn_add_user.clicked.connect(self.add_user)
        
        self.btn_update_user = QPushButton("Update User")
        self.btn_update_user.setStyleSheet(btn_style + "QPushButton{background-color: #6D28D9;} QPushButton:hover{background-color: #7C3AED;}")
        self.btn_update_user.clicked.connect(self.update_user)
        
        self.btn_remove_user = QPushButton("Remove User")
        self.btn_remove_user.setStyleSheet(btn_style + "QPushButton{background-color: #EF4444;} QPushButton:hover{background-color: #DC2626;}")
        self.btn_remove_user.clicked.connect(self.remove_user)
        
        form_layout.addWidget(self.btn_add_user)
        form_layout.addWidget(self.btn_update_user)
        form_layout.addWidget(self.btn_remove_user)
        
        self.user_err_label = QLabel("")
        self.user_err_label.setStyleSheet("font-family: 'Poppins', sans-serif; color: #EF4444; font-weight: bold; font-size: 11px;")
        self.user_err_label.setWordWrap(True)
        form_layout.addWidget(self.user_err_label)
        
        table_panel = QWidget()
        table_layout = QVBoxLayout(table_panel)
        table_layout.setContentsMargins(0, 0, 0, 0)
        
        self.user_table = QTableWidget()
        self.user_table.setColumnCount(8)
        self.user_table.setHorizontalHeaderLabels(["Username", "Doc Name", "Password", "Reg No", "Hospital", "Spec", "Phone", "Email"])
        self.user_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.user_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.user_table.cellClicked.connect(self.fill_user_form)
        
        # High-End Database Table Stylesheet
        self.user_table.setStyleSheet("""
            QTableWidget {
                background-color: rgba(255, 255, 255, 0.85);
                gridline-color: rgba(109, 40, 217, 0.1);
                border: 1px solid rgba(109, 40, 217, 0.15);
                border-radius: 12px;
                font-family: 'Poppins', 'Inter', sans-serif;
                font-size: 12px;
                color: #2E1065;
            }
            QTableWidget::item {
                padding: 6px;
            }
            QHeaderView::section {
                background-color: rgba(109, 40, 217, 0.08);
                color: #4C1D95;
                font-family: 'Poppins', 'Inter', sans-serif;
                font-weight: bold;
                border: none;
                border-bottom: 1px solid rgba(109, 40, 217, 0.15);
                padding: 8px;
            }
        """)
        
        table_layout.addWidget(self.user_table)
        
        layout.addWidget(form_panel)
        layout.addWidget(table_panel)
        self.tab_widget.addTab(self.user_tab, "Users Database")

    def init_manager_tab(self):
        self.manager_tab = QWidget()
        layout = QHBoxLayout(self.manager_tab)
        layout.setContentsMargins(15, 15, 15, 15)
        
        # Frosted glass form card
        form_panel = QWidget()
        form_panel.setObjectName("ManagerForm")
        form_panel.setFixedWidth(340)
        form_panel.setStyleSheet("""
            #ManagerForm {
                background-color: rgba(255, 255, 255, 0.72);
                border: 1px solid rgba(109, 40, 217, 0.15);
                border-radius: 16px;
            }
        """)
        
        shadow_mgr_form = QGraphicsDropShadowEffect(self)
        shadow_mgr_form.setBlurRadius(25)
        shadow_mgr_form.setColor(QColor(109, 40, 217, 15))
        form_panel.setGraphicsEffect(shadow_mgr_form)
        
        form_layout = QVBoxLayout(form_panel)
        form_layout.setSpacing(12)
        form_layout.setContentsMargins(20, 20, 20, 20)
        
        self.m_username = QLineEdit()
        self.m_username.setPlaceholderText("Username")
        self.m_doctor_name = QLineEdit()
        self.m_doctor_name.setPlaceholderText("Doctor Name")
        self.m_reg_no = QLineEdit()
        self.m_reg_no.setPlaceholderText("Registration ID")
        
        input_style = """
            QLineEdit {
                border: 1px solid rgba(109, 40, 217, 0.18);
                border-radius: 6px;
                padding: 8px 12px;
                height : 30px;
                font-family: 'Poppins', 'Inter', sans-serif;
                font-size: 13px;
                background-color: rgba(255, 255, 255, 0.9);
                color: #2E1065;
            }
            QLineEdit:focus {
                border: 1.5px solid #6D28D9;
                background-color: #ffffff;
            }
        """
        self.m_username.setStyleSheet(input_style)
        self.m_doctor_name.setStyleSheet(input_style)
        self.m_reg_no.setStyleSheet(input_style)
        
        lbl_style = "font-family: 'Poppins', sans-serif; font-size: 11px; color: #4C1D95; font-weight: bold;"
        
        form_layout.addWidget(QLabel("Search and Promote User:", styleSheet="font-family: 'Poppins', sans-serif; font-size: 12px; color: #6D28D9; font-weight: bold;"))
        form_layout.addWidget(QLabel("Target Username:", styleSheet=lbl_style))
        form_layout.addWidget(self.m_username)
        form_layout.addWidget(QLabel("Target Doctor Name:", styleSheet=lbl_style))
        form_layout.addWidget(self.m_doctor_name)
        form_layout.addWidget(QLabel("Target Reg ID:", styleSheet=lbl_style))
        form_layout.addWidget(self.m_reg_no)
        
        btn_style = """
            QPushButton {
                font-family: 'Poppins', 'Inter', sans-serif;
                font-weight: bold;
                border: none;
                border-radius: 6px;
                padding: 11px;
                color: white;
                font-size: 13px;
                height: 34px;
            }
        """
        
        self.btn_add_manager = QPushButton("Add Manager")
        self.btn_add_manager.setStyleSheet(btn_style + "QPushButton{background-color: #10B981;} QPushButton:hover{background-color: #059669;}")
        self.btn_add_manager.clicked.connect(self.add_manager)
        
        self.btn_remove_manager = QPushButton("Remove Manager")
        self.btn_remove_manager.setStyleSheet(btn_style + "QPushButton{background-color: #EF4444;} QPushButton:hover{background-color: #DC2626;}")
        self.btn_remove_manager.clicked.connect(self.remove_manager)
        
        form_layout.addWidget(self.btn_add_manager)
        form_layout.addWidget(self.btn_remove_manager)
        
        self.mgr_err_label = QLabel("")
        self.mgr_err_label.setStyleSheet("font-family: 'Poppins', sans-serif; color: #EF4444; font-weight: bold; font-size: 11px;")
        self.mgr_err_label.setWordWrap(True)
        form_layout.addWidget(self.mgr_err_label)
        
        table_panel = QWidget()
        table_layout = QVBoxLayout(table_panel)
        table_layout.setContentsMargins(0, 0, 0, 0)
        
        self.manager_table = QTableWidget()
        self.manager_table.setColumnCount(2)
        self.manager_table.setHorizontalHeaderLabels(["Username", "Password"])
        self.manager_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.manager_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.manager_table.cellClicked.connect(self.fill_manager_form)
        self.manager_table.setStyleSheet("""
            QTableWidget {
                background-color: rgba(255, 255, 255, 0.85);
                gridline-color: rgba(109, 40, 217, 0.1);
                border: 1px solid rgba(109, 40, 217, 0.15);
                border-radius: 12px;
                font-family: 'Poppins', 'Inter', sans-serif;
                font-size: 12px;
                color: #2E1065;
            }
            QTableWidget::item {
                padding: 6px;
            }
            QHeaderView::section {
                background-color: rgba(109, 40, 217, 0.08);
                color: #4C1D95;
                font-family: 'Poppins', 'Inter', sans-serif;
                font-weight: bold;
                border: none;
                border-bottom: 1px solid rgba(109, 40, 217, 0.15);
                padding: 8px;
            }
        """)
        
        table_layout.addWidget(self.manager_table)
        
        layout.addWidget(form_panel)
        layout.addWidget(table_panel)
        self.tab_widget.addTab(self.manager_tab, "Managers Database")

    def validate_password(self, password):
        has_digit = any(char.isdigit() for char in password)
        has_special = any(not char.isalnum() for char in password)
        return has_digit and has_special

    def validate_phone(self, phone):
        return len(phone) == 10 and phone.isdigit()

    def validate_email(self, email):
        if email.endswith("@gmail.com"):
            return True
        pattern = r"^[^@]+@[^@]+\.(com|in|edu)$"
        return bool(re.match(pattern, email))

    def load_excel_data(self, filepath):
        data = []
        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)
            sheet = wb.active
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if any(x is not None for x in row):
                    data.append(list(row))
            wb.close()
        except Exception:
            pass
        return data

    def save_excel_data(self, filepath, headers, data):
        try:
            wb = openpyxl.Workbook()
            sheet = wb.active
            sheet.append(headers)
            for row in data:
                sheet.append(row)
            wb.save(filepath)
            wb.close()
        except Exception:
            pass

    def refresh_user_table(self):
        self.user_table.setRowCount(0)
        data = self.load_excel_data(self.user_file)
        
        for row_data in data:
            username = row_data[0]
            if username not in self.user_states:
                self.user_states[username] = "original"
                
        for row_idx, row_data in enumerate(data):
            self.user_table.insertRow(row_idx)
            username = row_data[0]
            state = self.user_states.get(username, "original")
            
            # Premium low-saturation pastel tints mapped to database highlight tracks
            if state == "added":
                bg_color = QColor("#E0F2FE")    # Soft Sky Blue
            elif state == "updated":
                bg_color = QColor("#F3E8FF")  # Soft Lavender
            else:
                bg_color = QColor("#FAF7FF")  # Clean Frost White
                
            for col_idx, value in enumerate(row_data):
                val_str = str(value) if value is not None else ""
                item = QTableWidgetItem(val_str)
                item.setBackground(bg_color)
                item.setForeground(QColor("#2E1065"))
                self.user_table.setItem(row_idx, col_idx, item)

    def refresh_manager_table(self):
        self.manager_table.setRowCount(0)
        data = self.load_excel_data(self.manager_file)
        
        for row_data in data:
            username = row_data[0]
            if username not in self.manager_states:
                self.manager_states[username] = "original"
                
        for row_idx, row_data in enumerate(data):
            self.manager_table.insertRow(row_idx)
            username = row_data[0]
            state = self.manager_states.get(username, "original")
            
            if state == "added":
                bg_color = QColor("#E0F2FE")
            else:
                bg_color = QColor("#FAF7FF")
                
            for col_idx, value in enumerate(row_data):
                val_str = str(value) if value is not None else ""
                item = QTableWidgetItem(val_str)
                item.setBackground(bg_color)
                item.setForeground(QColor("#2E1065"))
                self.manager_table.setItem(row_idx, col_idx, item)

    def fill_user_form(self, row, column):
        self.u_username.setText(self.user_table.item(row, 0).text())
        self.u_doctor_name.setText(self.user_table.item(row, 1).text())
        self.u_password.setText(self.user_table.item(row, 2).text())
        self.u_reg_no.setText(self.user_table.item(row, 3).text())
        self.u_hospital.setText(self.user_table.item(row, 4).text())
        
        spec = self.user_table.item(row, 5).text()
        index = self.u_specialization.findText(spec)
        if index >= 0:
            self.u_specialization.setCurrentIndex(index)
            
        self.u_phone.setText(self.user_table.item(row, 6).text())
        self.u_email.setText(self.user_table.item(row, 7).text())

    def fill_manager_form(self, row, column):
        self.m_username.setText(self.manager_table.item(row, 0).text())

    def clear_user_form(self):
        self.u_username.clear()
        self.u_doctor_name.clear()
        self.u_password.clear()
        self.u_reg_no.clear()
        self.u_hospital.clear()
        self.u_phone.clear()
        self.u_email.clear()
        self.user_err_label.clear()

    def add_user(self):
        self.user_err_label.clear()
        username = self.u_username.text().strip()
        doctor_name = self.u_doctor_name.text().strip()
        password = self.u_password.text().strip()
        reg_no = self.u_reg_no.text().strip()
        hospital = self.u_hospital.text().strip()
        specialization = self.u_specialization.currentText()
        phone = self.u_phone.text().strip()
        email = self.u_email.text().strip()
        
        if not all([username, doctor_name, password, reg_no, hospital, phone, email]):
            self.user_err_label.setText("Error: All fields are required.")
            return
            
        if not self.validate_password(password):
            self.user_err_label.setText("Error: Password must contain at least one digit and one special character.")
            return
            
        if not self.validate_phone(phone):
            self.user_err_label.setText("Error: Phone number must be exactly 10 digits.")
            return
            
        if not self.validate_email(email):
            self.user_err_label.setText("Error: Invalid email format.")
            return
            
        users = self.load_excel_data(self.user_file)
        for u in users:
            if u[0] == username:
                self.user_err_label.setText("Error: Username must be unique.")
                return
            if str(u[3]) == reg_no:
                self.user_err_label.setText("Error: Registration No. must be unique.")
                return
                
        new_user = [username, doctor_name, password, reg_no, hospital, specialization, phone, email]
        users.append(new_user)
        self.save_excel_data(self.user_file, ["username", "doctor_name", "password", "reg_no", "hospital", "specialization", "phone", "email"], users)
        
        self.user_states[username] = "added"
        self.refresh_user_table()
        self.clear_user_form()

    def update_user(self):
        self.user_err_label.clear()
        username = self.u_username.text().strip()
        doctor_name = self.u_doctor_name.text().strip()
        password = self.u_password.text().strip()
        reg_no = self.u_reg_no.text().strip()
        hospital = self.u_hospital.text().strip()
        specialization = self.u_specialization.currentText()
        phone = self.u_phone.text().strip()
        email = self.u_email.text().strip()
        
        if not all([username, doctor_name, password, reg_no, hospital, phone, email]):
            self.user_err_label.setText("Error: All fields are required.")
            return
            
        if not self.validate_password(password):
            self.user_err_label.setText("Error: Password must contain at least one digit and one special character.")
            return
            
        if not self.validate_phone(phone):
            self.user_err_label.setText("Error: Phone number must be exactly 10 digits.")
            return
            
        if not self.validate_email(email):
            self.user_err_label.setText("Error: Invalid email format.")
            return
            
        users = self.load_excel_data(self.user_file)
        
        target_idx = -1
        for idx, u in enumerate(users):
            if u[0] == username:
                target_idx = idx
                break
                
        if target_idx == -1:
            self.user_err_label.setText("Error: User to update not found (username is key).")
            return
            
        for idx, u in enumerate(users):
            if idx != target_idx:
                if str(u[3]) == reg_no:
                    self.user_err_label.setText("Error: Registration No. must be unique.")
                    return
                    
        old_username = users[target_idx][0]
        users[target_idx] = [username, doctor_name, password, reg_no, hospital, specialization, phone, email]
        self.save_excel_data(self.user_file, ["username", "doctor_name", "password", "reg_no", "hospital", "specialization", "phone", "email"], users)
        
        if old_username in self.user_states:
            self.user_states.pop(old_username)
        self.user_states[username] = "updated"
        
        managers = self.load_excel_data(self.manager_file)
        manager_updated = False
        for m in managers:
            if m[0] == old_username:
                m[0] = username
                m[1] = password
                manager_updated = True
                if old_username in self.manager_states:
                    self.manager_states.pop(old_username)
                self.manager_states[username] = "updated"
                
        if manager_updated:
            self.save_excel_data(self.manager_file, ["Username", "Password"], managers)
            self.refresh_manager_table()
            
        self.refresh_user_table()
        self.clear_user_form()

    def remove_user(self):
        self.user_err_label.clear()
        username = self.u_username.text().strip()
        
        if not username:
            self.user_err_label.setText("Error: Please specify the username of the user to remove.")
            return
            
        users = self.load_excel_data(self.user_file)
        original_len = len(users)
        users = [u for u in users if u[0] != username]
        
        if len(users) == original_len:
            self.user_err_label.setText("Error: User not found.")
            return
            
        self.save_excel_data(self.user_file, ["username", "doctor_name", "password", "reg_no", "hospital", "specialization", "phone", "email"], users)
        
        if username in self.user_states:
            self.user_states.pop(username)
            
        managers = self.load_excel_data(self.manager_file)
        managers = [m for m in managers if m[0] != username]
        self.save_excel_data(self.manager_file, ["Username", "Password"], managers)
        
        if username in self.manager_states:
            self.manager_states.pop(username)
            
        self.refresh_user_table()
        self.refresh_manager_table()
        self.clear_user_form()

    def add_manager(self):
        self.mgr_err_label.clear()
        username = self.m_username.text().strip()
        doctor_name = self.m_doctor_name.text().strip()
        reg_id = self.m_reg_no.text().strip()
        
        if not all([username, doctor_name, reg_id]):
            self.mgr_err_label.setText("Error: All search fields required.")
            return
            
        users = self.load_excel_data(self.user_file)
        matched_user = None
        for u in users:
            if u[0] == username and u[1] == doctor_name and str(u[3]) == reg_id:
                matched_user = u
                break
                
        if matched_user is None:
            self.mgr_err_label.setText("Error: No matching user found in users database.")
            return
            
        managers = self.load_excel_data(self.manager_file)
        for m in managers:
            if m[0] == username:
                self.mgr_err_label.setText("Error: This user is already a manager.")
                return
                
        managers.append([matched_user[0], matched_user[2]])
        self.save_excel_data(self.manager_file, ["Username", "Password"], managers)
        
        self.manager_states[username] = "added"
        self.refresh_manager_table()
        self.m_username.clear()
        self.m_doctor_name.clear()
        self.m_reg_no.clear()

    def remove_manager(self):
        self.mgr_err_label.clear()
        username = self.m_username.text().strip()
        
        if not username:
            self.mgr_err_label.setText("Error: Please provide a username to remove.")
            return
            
        managers = self.load_excel_data(self.manager_file)
        original_len = len(managers)
        managers = [m for m in managers if m[0] != username]
        
        if len(managers) == original_len:
            self.mgr_err_label.setText("Error: Manager not found.")
            return
            
        self.save_excel_data(self.manager_file, ["Username", "Password"], managers)
        
        if username in self.manager_states:
            self.manager_states.pop(username)
            
        self.refresh_manager_table()
        self.m_username.clear()
        self.m_doctor_name.clear()
        self.m_reg_no.clear()

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = AdminDashboard()
    window.showFullScreen()
    window.show()
    sys.exit(app.exec())