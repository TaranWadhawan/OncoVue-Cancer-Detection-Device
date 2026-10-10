import sys
import os
import subprocess
import openpyxl
from PyQt6.QtWidgets import (QApplication, QMainWindow, QLabel, QVBoxLayout, 
                             QHBoxLayout, QWidget, QLineEdit, QPushButton, 
                             QComboBox, QTableWidget, QTableWidgetItem, QHeaderView, QGraphicsDropShadowEffect)
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QColor

class ReportListWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("OncoVue - Patient Reports Directory")
        self.setFixedSize(1280, 720)
        
        self.data_file = r"data/data.xlsx"
        self.user_file = r"data/user.xlsx"
        
        # Central Widget & Background Setup (Soft Pastel Gradient Base)
        self.central_widget = QWidget()
        self.central_widget.setObjectName("CentralWidget")
        self.central_widget.setStyleSheet("""
            #CentralWidget {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:1, 
                    stop:0 #FAF7FF, 
                    stop:0.6 #F3E8FF, 
                    stop:1 #E9D5FF);
            }
        """)
        self.setCentralWidget(self.central_widget)
        
        # Master Vertical Layout
        self.main_layout = QVBoxLayout(self.central_widget)
        self.main_layout.setContentsMargins(25, 20, 25, 20)
        self.main_layout.setSpacing(15)
        
        # Top Bar Layout (Glass Back Button)
        self.top_bar = QHBoxLayout()
        self.btn_back = QPushButton("← Back")
        self.btn_back.setFixedSize(110, 38)
        self.btn_back.setStyleSheet("""
            QPushButton {
                background-color: rgba(109, 40, 217, 0.08);
                border: 1px solid rgba(109, 40, 217, 0.18);
                color: #4C1D95; /* Velvet Slate */
                font-family: 'Poppins', 'Inter', sans-serif;
                font-weight: bold;
                border-radius: 8px;
                font-size: 13px;
            }
            QPushButton:hover {
                background-color: rgba(109, 40, 217, 0.15);
                border-color: rgba(109, 40, 217, 0.3);
            }
            QPushButton:pressed {
                background-color: rgba(109, 40, 217, 0.25);
            }
        """)
        self.btn_back.clicked.connect(self.back_to_dashboard)
        self.top_bar.addWidget(self.btn_back)
        self.top_bar.addStretch()
        self.main_layout.addLayout(self.top_bar)
        
        # Workspace split
        self.body_layout = QHBoxLayout()
        self.main_layout.addLayout(self.body_layout)
        
        # Left Filter Panel (Frosted Glass Container)
        self.filter_panel = QWidget()
        self.filter_panel.setObjectName("FilterPanel")
        self.filter_panel.setFixedWidth(290)
        self.filter_panel.setStyleSheet("""
            #FilterPanel {
                background-color: rgba(255, 255, 255, 0.72);
                border: 1px solid rgba(109, 40, 217, 0.15);
                border-radius: 16px;
            }
        """)
        
        shadow_filter = QGraphicsDropShadowEffect(self)
        shadow_filter.setBlurRadius(25)
        shadow_filter.setColor(QColor(109, 40, 217, 15))
        shadow_filter.setOffset(0, 8)
        self.filter_panel.setGraphicsEffect(shadow_filter)
        
        self.filter_layout = QVBoxLayout(self.filter_panel)
        self.filter_layout.setSpacing(12)
        self.filter_layout.setContentsMargins(25, 25, 25, 25)
        self.body_layout.addWidget(self.filter_panel)
        
        # Right Table Panel (Frosted Glass Container)
        self.table_panel = QWidget()
        self.table_panel.setObjectName("TablePanel")
        self.table_panel.setStyleSheet("""
            #TablePanel {
                background-color: rgba(255, 255, 255, 0.72);
                border: 1px solid rgba(109, 40, 217, 0.15);
                border-radius: 16px;
            }
        """)
        
        shadow_table = QGraphicsDropShadowEffect(self)
        shadow_table.setBlurRadius(25)
        shadow_table.setColor(QColor(109, 40, 217, 15))
        shadow_table.setOffset(0, 8)
        self.table_panel.setGraphicsEffect(shadow_table)
        
        self.table_layout = QVBoxLayout(self.table_panel)
        self.table_layout.setContentsMargins(20, 20, 20, 20)
        self.body_layout.addWidget(self.table_panel)
        
        self.setup_filter_panel()
        self.setup_table_panel()
        self.refresh_table()

    def back_to_dashboard(self):
                # self.timer.stop()
        subprocess.Popen([sys.executable, "loader.py"])
        subprocess.Popen([sys.executable, "page3.py"])
        self.close()

    def get_doctors_list(self):
        doctors = ["All Doctors"]
        if os.path.exists(self.user_file):
            try:
                wb = openpyxl.load_workbook(self.user_file, data_only=True)
                sheet = wb.active
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if len(row) >= 2 and row[1]:
                        doc_name = str(row[1]).strip()
                        if doc_name not in doctors:
                            doctors.append(doc_name)
                wb.close()
            except Exception:
                pass
        return doctors

    def setup_filter_panel(self):
        lbl_filter_title = QLabel("FILTER PATIENTS")
        lbl_filter_title.setStyleSheet("""
            font-family: 'Poppins', 'Inter', sans-serif;
            font-size: 11px;
            font-weight: bold;
            color: #6D28D9; /* Orchid Purple */
            letter-spacing: 2px;
            background: transparent;
        """)
        self.filter_layout.addWidget(lbl_filter_title)
        
        self.f_name = QLineEdit()
        self.f_name.setPlaceholderText("Search by Name")
        self.f_name.textChanged.connect(self.refresh_table)
        
        self.f_age = QLineEdit()
        self.f_age.setPlaceholderText("Filter by Age")
        self.f_age.textChanged.connect(self.refresh_table)
        
        self.f_gender = QComboBox()
        self.f_gender.addItems(["All Genders", "Female"])
        self.f_gender.currentIndexChanged.connect(self.refresh_table)
        
        self.f_doctor = QComboBox()
        self.f_doctor.addItems(self.get_doctors_list())
        self.f_doctor.currentIndexChanged.connect(self.refresh_table)
        
        self.f_hospital = QLineEdit()
        self.f_hospital.setPlaceholderText("Filter by Hospital")
        self.f_hospital.textChanged.connect(self.refresh_table)
        
        self.f_her = QComboBox()
        self.f_her.addItems(["All HER Scores", "0", "1", "2", "3"])
        self.f_her.currentIndexChanged.connect(self.refresh_table)
        
        input_style = """
            QLineEdit, QComboBox {
                border: 1px solid rgba(109, 40, 217, 0.18);
                border-radius: 6px;
                padding: 8px 12px;
                background-color: rgba(255, 255, 255, 0.9);
                color: #2E1065; /* Deep Amethyst */
                font-family: 'Poppins', 'Inter', sans-serif;
                font-size: 13px;
            }
            QLineEdit:focus, QComboBox:focus {
                border: 1.5px solid #6D28D9;
                background-color: #ffffff;
            }
        """
        self.f_name.setStyleSheet(input_style)
        self.f_age.setStyleSheet(input_style)
        self.f_gender.setStyleSheet(input_style)
        self.f_doctor.setStyleSheet(input_style)
        self.f_hospital.setStyleSheet(input_style)
        self.f_her.setStyleSheet(input_style)
        
        self.filter_layout.addWidget(self.f_name)
        self.filter_layout.addWidget(self.f_age)
        self.filter_layout.addWidget(self.f_gender)
        self.filter_layout.addWidget(self.f_doctor)
        self.filter_layout.addWidget(self.f_hospital)
        self.filter_layout.addWidget(self.f_her)
        
        # Custom Crimson Reset Action
        self.btn_reset = QPushButton("Reset Filters")
        self.btn_reset.setMinimumHeight(38)
        self.btn_reset.setStyleSheet("""
            QPushButton {
                background-color: #EF4444; /* Soft Crimson */
                color: white;
                font-family: 'Poppins', 'Inter', sans-serif;
                font-weight: bold;
                border: none;
                border-radius: 6px;
                font-size: 13px;
            }
            QPushButton:hover {
                background-color: #DC2626;
            }
        """)
        self.btn_reset.clicked.connect(self.reset_filters)
        self.filter_layout.addWidget(self.btn_reset)
        self.filter_layout.addStretch()

    def reset_filters(self):
        self.f_name.clear()
        self.f_age.clear()
        self.f_gender.setCurrentIndex(0)
        self.f_doctor.setCurrentIndex(0)
        self.f_hospital.clear()
        self.f_her.setCurrentIndex(0)
        self.refresh_table()

    def setup_table_panel(self):
        self.table = QTableWidget()
        self.table.setColumnCount(8)
        self.table.setHorizontalHeaderLabels(["Name", "Gender", "Age", "Phone", "Doctor", "Hospital", "HER Final", "Final Prediction"])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.cellDoubleClicked.connect(self.open_patient_report)
        
        # High-End Database Table Stylesheet
        self.table.setStyleSheet("""
            QTableWidget {
                background-color: rgba(255, 255, 255, 0.85);
                gridline-color: rgba(109, 40, 217, 0.1);
                border: none;
                border-radius: 12px;
                font-family: 'Poppins', 'Inter', sans-serif;
                font-size: 13px;
                color: #2E1065; /* Deep Amethyst */
            }
            QTableWidget::item {
                padding: 10px;
                background-color: transparent;
            }
            QTableWidget::item:selected {
                background-color: rgba(192, 132, 252, 0.25); /* Translucent Lavender highlight */
                color: #2E1065;
            }
            QHeaderView::section {
                background-color: rgba(109, 40, 217, 0.08); /* Frosted header segment */
                color: #4C1D95; /* Velvet Slate header text */
                font-family: 'Poppins', 'Inter', sans-serif;
                font-weight: bold;
                border: none;
                border-bottom: 1px solid rgba(109, 40, 217, 0.15);
                padding: 10px;
            }
        """)
        self.table_layout.addWidget(self.table)

    def load_excel_data(self):
        data = []
        if not os.path.exists(self.data_file):
            return data
        try:
            wb = openpyxl.load_workbook(self.data_file, data_only=True)
            sheet = wb.active
            headers = [str(cell.value).strip().lower() for cell in sheet[1]]
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if any(x is not None for x in row):
                    row_dict = {headers[i]: row[i] for i in range(len(headers)) if i < len(row)}
                    data.append(row_dict)
            wb.close()
        except Exception:
            pass
        return data

    def refresh_table(self):
        self.table.setRowCount(0)
        data = self.load_excel_data()
        
        name_q = self.f_name.text().strip().lower()
        age_q = self.f_age.text().strip().lower()
        gender_q = self.f_gender.currentText().strip().lower()
        doctor_q = self.f_doctor.currentText().strip().lower()
        hospital_q = self.f_hospital.text().strip().lower()
        her_q = self.f_her.currentText().strip().lower()
        
        filtered_data = []
        for row in data:
            p_name = str(row.get("patien_name", "")).strip().lower()
            p_age = str(row.get("age", "")).strip().lower()
            p_gender = str(row.get("gender", "")).strip().lower()
            p_doctor = str(row.get("doctor", "")).strip().lower()
            p_hospital = str(row.get("hospital", "")).strip().lower()
            p_her = str(row.get("her_final", "")).strip().lower()
            
            match = True
            if name_q and name_q not in p_name:
                match = False
            if age_q and age_q != p_age:
                match = False
            if gender_q != "all genders" and gender_q != p_gender:
                match = False
            if doctor_q != "all doctors" and doctor_q != p_doctor:
                match = False
            if hospital_q and hospital_q not in p_hospital:
                match = False
            if her_q != "all her scores" and her_q != p_her:
                match = False
                
            if match:
                filtered_data.append(row)
                
        for row_idx, row_data in enumerate(filtered_data):
            self.table.insertRow(row_idx)
            self.table.setItem(row_idx, 0, QTableWidgetItem(str(row_data.get("patien_name", ""))))
            self.table.setItem(row_idx, 1, QTableWidgetItem(str(row_data.get("gender", ""))))
            self.table.setItem(row_idx, 2, QTableWidgetItem(str(row_data.get("age", ""))))
            self.table.setItem(row_idx, 3, QTableWidgetItem(str(row_data.get("phone_no", ""))))
            self.table.setItem(row_idx, 4, QTableWidgetItem(str(row_data.get("doctor", ""))))
            self.table.setItem(row_idx, 5, QTableWidgetItem(str(row_data.get("hospital", ""))))
            self.table.setItem(row_idx, 6, QTableWidgetItem(str(row_data.get("her_final", ""))))
            self.table.setItem(row_idx, 7, QTableWidgetItem(str(row_data.get("final_prediction", ""))))

    def open_patient_report(self, row, column):
        phone_no = self.table.item(row, 3).text()
                # self.timer.stop()
        subprocess.Popen([sys.executable, "loader.py"])
        subprocess.Popen([sys.executable, "report_page2.py", phone_no])
        self.close()

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = ReportListWindow()
    window.showFullScreen()
    window.show()
    sys.exit(app.exec())