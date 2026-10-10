import sys
import os
import subprocess
import openpyxl
from PyQt6.QtWidgets import (QApplication, QMainWindow, QLabel, QVBoxLayout, 
                             QHBoxLayout, QWidget, QLineEdit, QPushButton, QCheckBox, QGraphicsDropShadowEffect)
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QColor
from keyboard import VirtualKeyboard 

class AdminLoginWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("OncoVue - Manager Verification")
        self.setFixedSize(1280, 720)
        
        # Central Widget & Background Setup (Soft Pastel Gradient Base)
        self.central_widget = QWidget()
        self.central_widget.setObjectName("CentralWidget")
        self.central_widget.setStyleSheet("""
            #CentralWidget {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:1, 
                    stop:0 #FAF7FF, 
                    stop:0.6 #F3E8FF, 
                    stop:1 #E9D5FF);
            }
        """)
        self.setCentralWidget(self.central_widget)
        
        # Master Vertical Layout
        self.main_layout = QVBoxLayout(self.central_widget)
        self.main_layout.setContentsMargins(25, 20, 25, 20)
        
        # Top Navigation Bar Layout (Glass Back Button)
        self.top_bar = QHBoxLayout()
        self.btn_back = QPushButton("← Back")
        self.btn_back.setFixedSize(110, 38)
        self.btn_back.setStyleSheet("""
            QPushButton {
                background-color: rgba(109, 40, 217, 0.08);
                border: 1px solid rgba(109, 40, 217, 0.18);
                color: #4C1D95; /* Velvet Slate */
                font-family: 'Poppins', 'Inter', sans-serif;
                font-weight: bold;
                border-radius: 8px;
                font-size: 13px;
            }
            QPushButton:hover {
                background-color: rgba(109, 40, 217, 0.15);
                border-color: rgba(109, 40, 217, 0.3);
            }
            QPushButton:pressed {
                background-color: rgba(109, 40, 217, 0.25);
            }
        """)
        self.btn_back.clicked.connect(self.back_to_page3)
        self.top_bar.addWidget(self.btn_back)
        self.top_bar.addStretch()
        self.main_layout.addLayout(self.top_bar)
        
        self.main_layout.addStretch()
        
        # Frosted glass verification card container
        self.card_widget = QWidget()
        self.card_widget.setObjectName("CardWidget")
        self.card_widget.setFixedSize(420, 480)
        self.card_widget.setStyleSheet("""
            #CardWidget {
                background-color: rgba(255, 255, 255, 0.7);
                border: 1px solid rgba(109, 40, 217, 0.15);
                border-radius: 20px;
            }
        """)
        
        # Soft outer depth shadow
        shadow = QGraphicsDropShadowEffect(self)
        shadow.setBlurRadius(30)
        shadow.setColor(QColor(109, 40, 217, 25))
        shadow.setOffset(0, 10)
        self.card_widget.setGraphicsEffect(shadow)
        
        self.main_layout.addWidget(self.card_widget, alignment=Qt.AlignmentFlag.AlignCenter)
        
        # Internal Card Layout
        self.card_layout = QVBoxLayout(self.card_widget)
        self.card_layout.setContentsMargins(40, 45, 40, 45)
        self.card_layout.setSpacing(18)
        
        # Small decor text label
        self.brand_label = QLabel("OncoVue SECURITY")
        self.brand_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.brand_label.setStyleSheet("""
            font-family: 'Poppins', 'Inter', sans-serif;
            font-size: 10px;
            font-weight: bold;
            color: #C084FC; /* Lavender Accent */
            letter-spacing: 4px;
            background: transparent;
        """)
        self.card_layout.addWidget(self.brand_label)
        
        # Section Header Title
        self.title_label = QLabel("Manager Login")
        self.title_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.title_label.setStyleSheet("""
            font-family: 'Cormorant Garamond', 'Georgia', serif;
            font-size: 32px;
            font-weight: 500;
            color: #2E1065; /* Deep Amethyst */
            background: transparent;
        """)
        self.card_layout.addWidget(self.title_label)
        
        # Form Input Fields
        input_style = """
            QLineEdit {
                border: 1px solid rgba(109, 40, 217, 0.2);
                border-radius: 8px;
                padding-left: 14px;
                font-family: 'Poppins', 'Inter', sans-serif;
                font-size: 13px;
                color: #2E1065;
                background-color: rgba(255, 255, 255, 0.9);
            }
            QLineEdit:focus {
                border: 1.5px solid #6D28D9;
                background-color: #ffffff;
            }
        """
        
        self.username_input = QLineEdit()
        self.username_input.setPlaceholderText("Manager Username")
        self.username_input.setMinimumHeight(45)
        self.username_input.setStyleSheet(input_style)
        self.card_layout.addWidget(self.username_input)
        
        self.password_input = QLineEdit()
        self.password_input.setPlaceholderText("Password")
        self.password_input.setEchoMode(QLineEdit.EchoMode.Password)
        self.password_input.setMinimumHeight(45)
        self.password_input.setStyleSheet(input_style)
        self.card_layout.addWidget(self.password_input)
        
        # Toggle Layout Section
        self.toggle_layout = QHBoxLayout()
        self.show_pass_checkbox = QCheckBox("Show Password")
        self.show_pass_checkbox.setStyleSheet("""
            QCheckBox {
                font-family: 'Poppins', 'Inter', sans-serif;
                font-size: 11px;
                color: #4C1D95; /* Velvet Slate */
                background: transparent;
            }
            QCheckBox::indicator {
                width: 14px;
                height: 14px;
                border: 1px solid rgba(109, 40, 217, 0.3);
                border-radius: 3px;
                background: white;
            }
            QCheckBox::indicator:checked {
                background-color: #EC4899;
                border-color: #EC4899;
            }
        """)
        self.show_pass_checkbox.stateChanged.connect(self.toggle_password_visibility)
        self.toggle_layout.addWidget(self.show_pass_checkbox)
        self.card_layout.addLayout(self.toggle_layout)
        
        # Gradient Call-to-Action Action Button
        self.login_button = QPushButton("Login")
        self.login_button.setMinimumHeight(45)
        self.login_button.setStyleSheet("""
            QPushButton {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0, 
                    stop:0 #EC4899, 
                    stop:1 #6D28D9); /* Rose to Orchid gradient */
                color: white;
                font-family: 'Poppins', 'Inter', sans-serif;
                font-size: 14px;
                font-weight: bold;
                border: none;
                border-radius: 8px;
            }
            QPushButton:hover {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0, 
                    stop:0 #F472B6, 
                    stop:1 #7C3AED);
            }
            QPushButton:pressed {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0, 
                    stop:0 #DB2777, 
                    stop:1 #5B21B6);
            }
        """)
        self.login_button.clicked.connect(self.handle_login)
        self.card_layout.addWidget(self.login_button)
        
        # Operation Status Label
        self.status_label = QLabel("")
        self.status_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.status_label.setStyleSheet("""
            font-family: 'Poppins', sans-serif;
            color: #EF4444; /* Crimson Red warn color */
            font-size: 12px;
            font-weight: 500;
            background: transparent;
        """)
        self.card_layout.addWidget(self.status_label)
        
        self.main_layout.addStretch()
        self.keyboard = VirtualKeyboard(self)
        self.keyboard.show()
        self.keyboard.move(100, 480)

    def back_to_page3(self):
        subprocess.Popen([sys.executable, "page3.py"])
        self.close()

    def toggle_password_visibility(self, state):
        if state == 2:
            self.password_input.setEchoMode(QLineEdit.EchoMode.Normal)
        else:
            self.password_input.setEchoMode(QLineEdit.EchoMode.Password)

    def create_default_excel(self, filename):
        try:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["Username", "Password"])
            sheet.append(["admin", "admin"])
            workbook.save(filename)
            workbook.close()
        except Exception:
            pass

    def validate_credentials(self, username, password):
        filename = r"data/manager.xlsx"
        
        if not os.path.exists(filename) or os.path.getsize(filename) == 0:
            self.create_default_excel(filename)
            
        try:
            workbook = openpyxl.load_workbook(filename, data_only=True)
        except Exception:
            self.create_default_excel(filename)
            try:
                workbook = openpyxl.load_workbook(filename, data_only=True)
            except Exception as e:
                self.status_label.setText(f"File Error: {str(e)}")
                return False
                
        try:
            sheet = workbook.active
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if len(row) >= 2:
                    db_user = str(row[0]).strip() if row[0] is not None else ""
                    db_pass = str(row[1]).strip() if row[1] is not None else ""
                    if db_user == username.strip() and db_pass == password.strip():
                        workbook.close()
                        return True
            workbook.close()
        except Exception as e:
            self.status_label.setText(f"Read Error: {str(e)}")
            
        return False

    def handle_login(self):
        username = self.username_input.text()
        password = self.password_input.text()
        
        if not username or not password:
            self.status_label.setText("Please fill in all fields")
            return
            
        if self.validate_credentials(username, password):
            # subprocess.Popen([sys.executable, "manage_admin_page2.py"])
            self.close()
        else:
            if "Error" not in self.status_label.text():
                self.status_label.setText("Invalid manager credentials")

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = AdminLoginWindow()
    window.show()
    sys.exit(app.exec())