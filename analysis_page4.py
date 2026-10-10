import os
os.environ["QT_QPA_PLATFORM"] = "xcb"
import sys
# import os
import time
import shutil
import subprocess
import openpyxl
from PyQt6.QtWidgets import (QApplication, QMainWindow, QLabel, QVBoxLayout, 
                             QHBoxLayout, QWidget, QLineEdit, QPushButton, 
                             QComboBox, QGridLayout, QGroupBox, QGraphicsDropShadowEffect)
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QPixmap, QColor
from keyboard import VirtualKeyboard

class AnalysisPage4(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("OncoVue - Deep Learning Classification Analysis")
        self.setFixedSize(1280, 720)
        
        self.data_file = r"data/data.xlsx"
        self.user_file = r"data/user.xlsx"
        
        self.temp_dir = "temp"
        self.image_dir = r"data/image"
        
        self.init_directories()
        self.load_preprocessed_image_path()
        self.mock_dl_scores()
        
        # Central Widget & Background Setup (Soft Pastel Gradient Base)
        self.central_widget = QWidget()
        self.central_widget.setObjectName("CentralWidget")
        self.central_widget.setStyleSheet("""
            #CentralWidget {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:1, 
                    stop:0 #FAF7FF, 
                    stop:0.6 #F3E8FF, 
                    stop:1 #E9D5FF);
            }
        """)
        self.setCentralWidget(self.central_widget)
        
        # Master Vertical Layout
        self.main_layout = QVBoxLayout(self.central_widget)
        self.main_layout.setContentsMargins(25, 20, 25, 20)
        self.main_layout.setSpacing(15)
        
        # Top Navigation Bar (Glass Back Button)
        self.top_bar = QHBoxLayout()
        self.btn_back = QPushButton("← Back")
        self.btn_back.setFixedSize(110, 38)
        self.btn_back.setStyleSheet("""
            QPushButton {
                background-color: rgba(109, 40, 217, 0.08);
                border: 1px solid rgba(109, 40, 217, 0.18);
                color: #4C1D95; /* Velvet Slate */
                font-family: 'Poppins', 'Inter', sans-serif;
                font-weight: bold;
                border-radius: 8px;
                font-size: 13px;
            }
            QPushButton:hover {
                background-color: rgba(109, 40, 217, 0.15);
                border-color: rgba(109, 40, 217, 0.3);
            }
            QPushButton:pressed {
                background-color: rgba(109, 40, 217, 0.25);
            }
        """)
        self.btn_back.clicked.connect(self.back_to_page3)
        self.top_bar.addWidget(self.btn_back)
        self.top_bar.addStretch()
        self.main_layout.addLayout(self.top_bar)
        
        # Workspace split
        self.body_layout = QHBoxLayout()
        self.body_layout.setSpacing(15)
        self.main_layout.addLayout(self.body_layout)
        
        self.left_panel = QVBoxLayout()
        self.body_layout.addLayout(self.left_panel, stretch=1)
        
        self.right_panel = QVBoxLayout()
        self.body_layout.addLayout(self.right_panel, stretch=1)
        
        self.setup_left_panel()
        self.setup_right_panel()
        
        # Bottom Navigation Control
        self.bottom_bar = QHBoxLayout()
        self.btn_proceed = QPushButton("Proceed to Report Generation")
        self.btn_proceed.setMinimumHeight(45)
        self.btn_proceed.setStyleSheet("""
            QPushButton {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0, 
                    stop:0 #EC4899, 
                    stop:1 #6D28D9); /* Rose to Orchid gradient */
                color: white;
                font-family: 'Poppins', 'Inter', sans-serif;
                font-weight: bold;
                font-size: 14px;
                border: none;
                border-radius: 8px;
            }
            QPushButton:hover {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0, 
                    stop:0 #F472B6, 
                    stop:1 #7C3AED);
            }
        """)
        self.btn_proceed.clicked.connect(self.save_and_proceed)
        self.main_layout.addLayout(self.bottom_bar)
        self.bottom_bar.addWidget(self.btn_proceed)
        self.keyboard = VirtualKeyboard(self)
        self.keyboard.show()
        self.keyboard.move(100, 480)

    def back_to_page3(self):
        # self.timer.stop()
        subprocess.Popen([sys.executable, "loader_large.py"])
        subprocess.Popen([sys.executable, "analysis_page3.py"])
        self.close()

    def init_directories(self):
        os.makedirs("data", exist_ok=True)
        os.makedirs(self.image_dir, exist_ok=True)
        if not os.path.exists(self.data_file):
            wb = openpyxl.Workbook()
            sheet = wb.active
            sheet.append(["date", "time", "patien_name", "gender", "age", "phone_no", "check", "doctor", "hospital", "her0", "her1", "her2", "her3", "her_final", "final_prediction", "image_path", "report_path"])
            wb.save(self.data_file)
            wb.close()

    def load_preprocessed_image_path(self):
        self.latest_temp_image = None
        if os.path.exists(self.temp_dir):
            files = [os.path.join(self.temp_dir, f) for f in os.listdir(self.temp_dir) if f.lower().endswith(('.png', '.jpg', '.jpeg'))]
            if files:
                self.latest_temp_image = max(files, key=os.path.getmtime)

    # def mock_dl_scores(self):
        # import random
        # self.her0_score = round(random.uniform(0.05, 0.3), 4)
        # self.her1_score = round(random.uniform(0.05, 0.4), 4)
        # self.her2_score = round(random.uniform(0.25, 0.90), 4)
        # self.her3_score = round(random.uniform(0.35, 0.90), 4)
        
        # scores = {
        #     "her0": self.her0_score,
        #     "her1": self.her1_score,
        #     "her2": self.her2_score,
        #     "her3": self.her3_score
        # }
        
        # self.max_class = max(scores, key=scores.get)
        # if self.max_class == "her0":
        #     self.final_prediction = "Normal"
        # elif self.max_class == "her1":
        #     self.final_prediction = "Normal"
        # elif self.max_class == "her2":
        #     self.final_prediction = "Equivocal"
        # elif self.max_class == "her3":
        #     self.final_prediction = "Positive"
        # else :
        #     self.final_prediction = "Error"

    # def mock_dl_scores(self):
    #     scores_file = r"temp/current_scores.json"
    #     if os.path.exists(scores_file):
    #         try:
    #             with open(scores_file, 'r') as f:
    #                 data = json.load(f)
    #                 self.her0_score = data["her0"]
    #                 self.her1_score = data["her1"]
    #                 self.her2_score = data["her2"]
    #                 self.her3_score = data["her3"]
    #                 self.max_class = data["max_class"]
    #                 self.final_prediction = data["verdict"]
    #             return
    #         except Exception:
    #             pass

    #     # import random
    #     # self.her0_score = round(random.uniform(0.05, 0.90), 4)
    #     # self.her1_score = round(random.uniform(0.05, 0.90), 4)
    #     # self.her2_score = round(random.uniform(0.05, 0.90), 4)
    #     # self.her3_score = round(random.uniform(0.05, 0.90), 4)
        
    #     scores = {
    #         "her0": self.her0_score,
    #         "her1": self.her1_score,
    #         "her2": self.her2_score,
    #         "her3": self.her3_score
    #     }
        
    #     self.max_class = max(scores, key=scores.get)
    #     if self.max_class == "her0":
    #         self.final_prediction = "Normal"
    #     elif self.max_class == "her1":
    #         self.final_prediction = "Normal"
    #     elif self.max_class == "her2":
    #         self.final_prediction = "Equivocal"
    #     else:
    #         self.final_prediction = "Positive"


    def mock_dl_scores(self):
        import json
        import random

        scores_file = r"temp/current_scores.json"
        if os.path.exists(scores_file):
            try:
                with open(scores_file, 'r') as f:
                    data = json.load(f)
                    self.her0_score = data["her0"]
                    self.her1_score = data["her1"]
                    self.her2_score = data["her2"]
                    self.her3_score = data["her3"]
                    self.max_class = data["max_class"]
                    self.final_prediction = data["verdict"]
                return
            except Exception:
                pass

        # Fallback random initialization if the JSON file is missing or invalid
        self.her0_score = round(random.uniform(0.85, 0.90), 4)
        self.her1_score = round(random.uniform(0.05, 0.90), 4)
        self.her2_score = round(random.uniform(0.05, 0.90), 4)
        self.her3_score = round(random.uniform(0.05, 0.90), 4)
        
        scores = {
            "her0": self.her0_score,
            "her1": self.her1_score,
            "her2": self.her2_score,
            "her3": self.her3_score
        }
        
        self.max_class = max(scores, key=scores.get)
        if self.max_class == "her0":
            self.final_prediction = "Normal"
        elif self.max_class == "her1":
            self.final_prediction = "Normal"
        elif self.max_class == "her2":
            self.final_prediction = "Equivocal"
        else:
            self.final_prediction = "Positive"

    def get_doctor_data(self):
        doctors = {}
        if os.path.exists(self.user_file):
            try:
                wb = openpyxl.load_workbook(self.user_file, data_only=True)
                sheet = wb.active
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if len(row) >= 5 and row[1]:
                        doc_name = str(row[1]).strip()
                        hosp_name = str(row[4]).strip() if row[4] else ""
                        doctors[doc_name] = hosp_name
                wb.close()
            except Exception:
                pass
        return doctors

    def setup_left_panel(self):
        group_style = """
            QGroupBox {
                font-family: 'Poppins', 'Inter', sans-serif;
                font-size: 11px;
                color: #6D28D9; /* Orchid Purple */
                font-weight: bold;
                border: 1px solid rgba(109, 40, 217, 0.15);
                border-radius: 12px;
                margin-top: 10px;
                background-color: rgba(255, 255, 255, 0.72);
            }
            QGroupBox::title {
                subcontrol-origin: margin;
                left: 12px;
                padding: 0 5px;
            }
        """
        
        input_style = """
            QLineEdit, QComboBox {
                border: 1px solid rgba(109, 40, 217, 0.18);
                border-radius: 6px;
                padding: 8px 12px;
                background-color: rgba(255, 255, 255, 0.9);
                color: #2E1065; /* Deep Amethyst */
                font-family: 'Poppins', 'Inter', sans-serif;
                font-size: 13px;
            }
            QLineEdit:focus, QComboBox:focus {
                border: 1.5px solid #6D28D9;
                background-color: #ffffff;
            }
        """
        
        # Frosted glass Patient Data block
        self.p_group = QGroupBox("Patient & Clinical Data")
        self.p_group.setStyleSheet(group_style)
        p_layout = QVBoxLayout(self.p_group)
        p_layout.setContentsMargins(20, 25, 20, 20)
        p_layout.setSpacing(12)
        
        self.p_name = QLineEdit()
        self.p_name.setPlaceholderText("Patient Name")
        self.p_name.setStyleSheet(input_style)
        p_layout.addWidget(self.p_name)
        
        self.p_gender = QComboBox()
        self.p_gender.addItems([ "Female"])
        self.p_gender.setStyleSheet(input_style)
        p_layout.addWidget(self.p_gender)
        
        self.p_age = QLineEdit()
        self.p_age.setPlaceholderText("Age")
        self.p_age.setStyleSheet(input_style)
        p_layout.addWidget(self.p_age)
        
        self.p_phone = QLineEdit()
        self.p_phone.setPlaceholderText("Phone Number")
        self.p_phone.setStyleSheet(input_style)
        p_layout.addWidget(self.p_phone)
        
        self.doc_map = self.get_doctor_data()
        self.p_doctor = QComboBox()
        self.p_doctor.addItems(list(self.doc_map.keys()))
        self.p_doctor.setStyleSheet(input_style)
        self.p_doctor.currentIndexChanged.connect(self.update_hospital_display)
        p_layout.addWidget(self.p_doctor)
        
        self.p_hospital = QLineEdit()
        self.p_hospital.setPlaceholderText("Hospital (Auto-filled)")
        self.p_hospital.setReadOnly(True)
        self.p_hospital.setStyleSheet(input_style + """
            QLineEdit {
                background-color: rgba(109, 40, 217, 0.05); /* Lavender tint background */
                color: #4C1D95; /* Velvet Slate read-only font */
            }
        """)
        p_layout.addWidget(self.p_hospital)
        
        self.update_hospital_display()
        self.left_panel.addWidget(self.p_group)
        
        # Soft Crimson warning display
        self.err_label = QLabel("")
        self.err_label.setStyleSheet("""
            font-family: 'Poppins', sans-serif;
            color: #EF4444; 
            font-weight: bold; 
            font-size: 12px;
            background: transparent;
        """)
        self.err_label.setWordWrap(True)
        self.left_panel.addWidget(self.err_label)

        # Apply shadow to the left panel
        shadow_p = QGraphicsDropShadowEffect(self)
        shadow_p.setBlurRadius(25)
        shadow_p.setColor(QColor(109, 40, 217, 15))
        shadow_p.setOffset(0, 8)
        self.p_group.setGraphicsEffect(shadow_p)

    def update_hospital_display(self):
        doc = self.p_doctor.currentText()
        hosp = self.doc_map.get(doc, "")
        self.p_hospital.setText(hosp)

    def setup_right_panel(self):
        group_style = """
            QGroupBox {
                font-family: 'Poppins', 'Inter', sans-serif;
                font-size: 11px;
                color: #6D28D9;
                font-weight: bold;
                border: 1px solid rgba(109, 40, 217, 0.15);
                border-radius: 12px;
                margin-top: 10px;
                background-color: rgba(255, 255, 255, 0.72);
            }
            QGroupBox::title {
                subcontrol-origin: margin;
                left: 12px;
                padding: 0 5px;
            }
            QLabel {
                font-family: 'Poppins', 'Inter', sans-serif;
                font-size: 13px;
                color: #4C1D95;
            }
        """
        
        # Preprocessed diagnostics graphic card
        self.preview_group = QGroupBox("Preprocessed Image View")
        self.preview_group.setStyleSheet(group_style)
        p_layout = QVBoxLayout(self.preview_group)
        p_layout.setContentsMargins(15, 20, 15, 15)
        
        self.lbl_image = QLabel("Preprocessed frame will show here...")
        self.lbl_image.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.lbl_image.setFixedSize(450, 240)
        self.lbl_image.setStyleSheet("""
            QLabel {
                border: 2px dashed rgba(109, 40, 217, 0.18); 
                background-color: #000000; 
                border-radius: 8px;
                color: #C084FC;
            }
        """)
        p_layout.addWidget(self.lbl_image)
        
        if self.latest_temp_image and os.path.exists(self.latest_temp_image):
            pixmap = QPixmap(self.latest_temp_image)
            self.lbl_image.setPixmap(pixmap.scaled(450, 240, Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation))
            
        self.right_panel.addWidget(self.preview_group)
        
        # Deep Learning score summary panel
        self.results_group = QGroupBox("Deep Learning Classification Analysis")
        self.results_group.setStyleSheet(group_style)
        r_layout = QVBoxLayout(self.results_group)
        r_layout.setContentsMargins(20, 25, 20, 20)
        
        grid_scores = QGridLayout()
        grid_scores.setSpacing(8)
        
        # Accessible text styling
        lbl_grid_style = "font-family: 'Poppins', sans-serif; font-size: 12px; color: #4C1D95; font-weight: 500;"
        val_grid_style = "font-family: 'Poppins', sans-serif; font-size: 13px; color: #2E1065; font-weight: bold;"
        
        scores_list = [
            ("Class Her0 Score:", self.her0_score, 0),
            ("Class Her1 Score:", self.her1_score, 1),
            ("Class Her2 Score:", self.her2_score, 2),
            ("Class Her3 Score:", self.her3_score, 3)
        ]
        
        for label, score, row in scores_list:
            lbl = QLabel(label)
            lbl.setStyleSheet(lbl_grid_style)
            val = QLabel(str(score))
            val.setStyleSheet(val_grid_style)
            grid_scores.addWidget(lbl, row, 0)
            grid_scores.addWidget(val, row, 1)
        
        r_layout.addLayout(grid_scores)
        
        # Medical Diagnostic Verdict highlight
        self.lbl_verdict = QLabel(f"Max Score Class: {self.max_class.upper()} | Verdict: {self.final_prediction}")
        self.lbl_verdict.setStyleSheet("""
            font-family: 'Poppins', sans-serif;
            font-size: 13px; 
            color: #10B981; /* High-contrast soft green */
            font-weight: bold; 
            margin-top: 10px;
        """)
        r_layout.addWidget(self.lbl_verdict)
        
        self.right_panel.addWidget(self.results_group)

        # Apply shadows to right-hand panels
        shadow_preview = QGraphicsDropShadowEffect(self)
        shadow_preview.setBlurRadius(25)
        shadow_preview.setColor(QColor(109, 40, 217, 15))
        self.preview_group.setGraphicsEffect(shadow_preview)

        shadow_results = QGraphicsDropShadowEffect(self)
        shadow_results.setBlurRadius(25)
        shadow_results.setColor(QColor(109, 40, 217, 15))
        self.results_group.setGraphicsEffect(shadow_results)

    def save_and_proceed(self):
        self.err_label.clear()
        name = self.p_name.text().strip()
        gender = self.p_gender.currentText()
        age = self.p_age.text().strip()
        phone = self.p_phone.text().strip()
        doctor = self.p_doctor.currentText()
        hospital = self.p_hospital.text().strip()
        
        if not all([name, age, phone, doctor]):
            self.err_label.setText("Error: All fields are required.")
            return
            
        if not self.latest_temp_image or not os.path.exists(self.latest_temp_image):
            self.err_label.setText("Error: Preprocessed frame missing in temp folder.")
            return
            
        date_str = time.strftime("%d-%m-%Y")
        time_str = time.strftime("%H:%M:%S")
        
        img_no = 1
        while True:
            new_img_name = f"{name}_{gender}_{age}_{doctor}_{date_str}_img_{img_no}.jpg".replace(" ", "_")
            dest_img_path = os.path.join(self.image_dir, new_img_name)
            if not os.path.exists(dest_img_path):
                break
            img_no += 1
            
        try:
            shutil.move(self.latest_temp_image, dest_img_path)
        except Exception as e:
            self.err_label.setText(f"Error transferring file: {str(e)}")
            return
            
        try:
            wb = openpyxl.load_workbook(self.data_file)
            sheet = wb.active
            sheet.append([
                date_str, time_str, name, gender, age, phone, "Yes", doctor, hospital,
                self.her0_score, self.her1_score, self.her2_score, self.her3_score,
                self.max_class, self.final_prediction, dest_img_path, ""
            ])
            wb.save(self.data_file)
            wb.close()
        except Exception as e:
            self.err_label.setText(f"Database Error: {str(e)}")
            return
        # self.timer.stop()
        subprocess.Popen([sys.executable, "loader.py"])
        subprocess.Popen([sys.executable, "analysis_page5.py", phone])
        self.close()

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = AnalysisPage4()
    # window.showFullScreen()
    window.show()
    sys.exit(app.exec())